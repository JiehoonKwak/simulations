"""Read official NHIS/HIRA yearbook workbooks; never alter source files."""
from pathlib import Path
from decimal import Decimal
import hashlib
import json
import openpyxl
from zipfile import ZipFile

PROJECT = Path(__file__).resolve().parents[1]
RAW = PROJECT / "data" / "raw" / "workforce"
OUT = PROJECT / "data" / "processed" / "workforce.json"
BASE = "https://www.nhis.or.kr/nhis/together/wbhaec06300m01.do"
DOWNLOADS = {
    "nhis-yearbook-2023-tables.zip": (10847858, 364634, "9f20b4e42c45e815256e7c6377a9e279d6e6934e59be36ac60416e94fc4bed1b"),
    "nhis-yearbook-2023.pdf": (10847858, 364636, "1b3c23af2cf09906c450e19e9447c942c42d7fddee8b11616a47906f32361181"),
    "nhis-yearbook-2024-tables.zip": (11007426, 365189, "fd9e7313b08b1a44a49ec6655df27efa142fb167147f639a89fc06af1cd56de6"),
    "nhis-yearbook-2024.pdf": (11007426, 365191, "adf65c36b2f3bdc8dabab0e9b9c49ed339d156c1824b0b6e8f127eda4edc3b52"),
}
WORKBOOKS = {
    "hira-providers-2023.xlsx": ("nhis-yearbook-2023-tables.zip", "01-2", "8e99ee8479ac0da2088abd4fb0aa3ae20bde6d30b12761103e1a30d2f7bb6955"),
    "hira-providers-2024.xlsx": ("nhis-yearbook-2024-tables.zip", "01-2", "5c2189741886bc7fba458baf60cde8bedcbbaea3b28ae819dfba99b6069caa99"),
    "nhis-care-2024.xlsx": ("nhis-yearbook-2024-tables.zip", "03-1", "bf3ad3e9038957eae571b28039f318322a8d6dea03a7bfbe5f924ce87f289506"),
}
PROVIDERS = [
    ("tertiary_hospital", "상급종합병원"),
    ("general_hospital", "종합병원"),
    ("hospital", "병원"),
    ("long_term_care_hospital", "요양병원"),
    ("mental_health_hospital", "정신병원"),
    ("clinic", "의원"),
    ("dental_hospital", "치과병원"),
    ("dental_clinic", "치과의원"),
    ("midwifery_clinic", "조산원"),
    ("hospitalized_health_center", "보건의료원"),
    ("health_center", "보건소"),
    ("health_subcenter", "보건지소"),
    ("primary_health_care_post", "보건진료소"),
    ("korean_medicine_hospital", "한방병원"),
    ("korean_medicine_clinic", "한의원"),
    ("pharmacy", "약국"),
]
REGIONS = ["seoul", "busan", "daegu", "incheon", "gwangju", "daejeon", "ulsan", "sejong", "gyeonggi", "gangwon", "chungbuk", "chungnam", "jeonbuk", "jeonnam", "gyeongbuk", "gyeongnam", "jeju"]
QUALIFICATIONS = ["physicians", "general_practitioners", "interns", "residents", "specialists"]
FORMS = ["national", "public", "school_foundation", "special_juridical_corporation", "religious_foundation", "social_welfare_foundation", "corporation_aggregate", "juridical_foundation", "company", "medical_corporation", "consumer_cooperative", "social_cooperative", "individual", "military_hospital", "other"]


def file_info(name, url, **extra):
    data = (RAW / name).read_bytes()
    return {"id": name, "local_path": "data/raw/workforce/" + name, "url": url, "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest(), **extra}


def loc(file, sheet, cells, year, page):
    offset = 82 if year == 2024 else 80
    pages = [page, page + 1] if page in [40, 42, 144] else [page]
    return {"file_id": file, "sheet": sheet, "cells": cells, "pdf_file_id": f"nhis-yearbook-{year}.pdf", "pdf_printed_pages": pages, "pdf_pages_1based": [p + offset for p in pages]}


def original_member(archive, prefix):
    with ZipFile(RAW / archive) as zipped:
        member = next(i for i in zipped.infolist() if i.filename.startswith(prefix) and i.filename.endswith(".xlsx"))
        return member.filename if member.flag_bits & 0x800 else member.filename.encode("cp437").decode("cp949")


def integer(sheet, address):
    value = sheet[address].value
    assert isinstance(value, (int, float)) and int(value) == value, (sheet.title, address, value)
    return int(value)


missing = [f"  {RAW / name}\n    {BASE}?mode=download&articleNo={article}&attachNo={attachment}" for name, (article, attachment, _) in DOWNLOADS.items() if not (RAW / name).is_file()]
if missing:
    raise SystemExit("Missing raw workforce inputs. Download these official files with the indicated local names:\n" + "\n".join(missing))
for name, (_, _, expected_hash) in DOWNLOADS.items():
    actual_hash = hashlib.sha256((RAW / name).read_bytes()).hexdigest()
    if actual_hash != expected_hash:
        raise SystemExit(f"Source hash mismatch for {name}: expected {expected_hash}, got {actual_hash}. Review the upstream revision before updating this extraction; output was not changed.")
for name, (archive, prefix, expected_hash) in WORKBOOKS.items():
    with ZipFile(RAW / archive) as zipped:
        members = [i for i in zipped.infolist() if i.filename.startswith(prefix) and i.filename.endswith(".xlsx")]
        if len(members) != 1:
            raise SystemExit(f"Expected one {prefix} workbook in {archive}; found {len(members)}.")
        workbook_bytes = zipped.read(members[0])
    if hashlib.sha256(workbook_bytes).hexdigest() != expected_hash:
        raise SystemExit(f"Archived workbook hash mismatch: {name}. Output was not changed.")
    if (RAW / name).exists() and (RAW / name).read_bytes() != workbook_bytes:
        raise SystemExit(f"Extracted workbook differs from the verified archive: {RAW / name}. Remove the altered local copy and rerun; output was not changed.")
    (RAW / name).write_bytes(workbook_bytes)


data = {
    "schema_version": 1,
    "title": "Korean NHI-reported physicians, institutions and care, 2023–2024",
    "retrieved_on": "2026-10-02",
    "observation_years": [2023, 2024],
    "simulation_start_year": 2026,
    "status": "Extracted empirical aggregates; no 2026 nowcast or joint physician microdata reconstruction.",
    "publishers": ["National Health Insurance Service", "Health Insurance Review & Assessment Service"],
    "statistics_approval_number": "920006",
    "sources": [],
    "license": {
        "2024_work_marking": "KOGL Type 1 (attribution)",
        "marking_url": "https://www.hira.or.kr/bbsDummy.do?brdBltNo=2431&brdScnBltNo=4&pgmid=HIRAA020045010000",
        "terms_url": "https://www.kogl.or.kr/info/licenseType1.do",
        "note": "Marking directly observed on the alternate HIRA 2024 yearbook post. NHIS attachment page and the first HIRA post carry no observed work-specific marking. 2023 file-specific marking was not verified. Raw files remain git-ignored; processed file contains selected factual aggregates with attribution, not report facsimiles.",
    },
    "revision_status": {
        "hira_corrected_post_date": "2026-01-13",
        "hira_corrected_post_url": "https://www.hira.or.kr/bbsDummy.do?brdBltNo=2322&brdScnBltNo=4&pgmid=HIRAA020045020000",
        "checked_against_hira_corrigendum": False,
        "note": "NHIS current PDF and workbooks were acquired directly. HIRA lists a corrected PDF and corrigendum; its download helper repeatedly timed out, so equality to that corrected attachment has not been established. Acquired NHIS PDF matches the SMHDB mirror byte-for-byte.",
        "mirror_download_url": "https://seoulmentalhealth.kr/comm/getFile?fileNo=1&fileTy=ATTACH&upperNo=2100",
    },
    "definitions": {
        "workforce_reference": "Year-end personnel reported by NHI-claiming providers to HIRA; physicians use an exclusive-contract (전속) criterion. Persons, not FTE, licensed totals, all employed physicians, or owners.",
        "physicians": "의사 only, excluding dentists and Korean medicine physicians; includes general practitioners, interns, residents and specialists.",
        "institution_reference": "Year-end NHI provider registrations, not all Medical Service Act registrations. Provider types remain separate legal/reporting categories.",
        "region": "Provider location; institution × region is observed. Physician region and institution margins are separate tables, not a joint distribution.",
        "individual_establishment": "개인 in an institution registration table; counts institutions, not owner physicians or self-employed clinicians.",
        "care_reference": "NHIS Table III-4, service-date 2024 care, reflecting payments January 2024 through April 2025. NHI insured service users; not all medical spending.",
        "visit_days": "입내원일수: patient visit-days or inpatient bed-days recorded in claims. Not unique patients, completed cases or quality-adjusted care.",
        "reimbursed_days": "요양급여일수: includes in-facility medication days; an overlapping visit/medication day counts once.",
        "patients": "Within-category unique patients. Patients can overlap across settings and provider types; do not sum them to a national unique-patient count.",
        "medical_expenses_thousand_krw": "Insurer payment plus statutory patient copayment for covered care, in KRW 1,000; not physician income or net institution profit. Workbook precision retained as decimal strings.",
        "insurer_benefits_thousand_krw": "NHIS share of covered expenses, in KRW 1,000; workbook precision retained as decimal strings.",
        "zeros": "Source workbook numeric zeros are retained. The PDF commonly displays them as a dash.",
    },
    "provider_types": [{"id": k, "label_ko": v} for k, v in PROVIDERS],
    "institutions_by_type": [],
    "institutions_by_region_and_type": [],
    "physicians_by_type": [],
    "physicians_by_region": [],
    "institutions_by_establishment_2024": [],
    "care_by_type_2024": [],
}

for year, article, archive_id, pdf_id, publication in [(2023, 10847858, 364634, 364636, "2024-11"), (2024, 11007426, 365189, 365191, "2025-11")]:
    archive_name = f"nhis-yearbook-{year}-tables.zip"
    data["sources"].append(file_info(archive_name, f"{BASE}?mode=download&articleNo={article}&attachNo={archive_id}", publication_month=publication, landing_url=f"{BASE}?mode=view&articleNo={article}"))
    data["sources"].append(file_info(f"nhis-yearbook-{year}.pdf", f"{BASE}?mode=download&articleNo={article}&attachNo={pdf_id}", publication_month=publication))
    name = f"hira-providers-{year}.xlsx"
    data["sources"].append(file_info(name, f"{BASE}?mode=download&articleNo={article}&attachNo={archive_id}", archive_file_id=archive_name, archive_member=original_member(archive_name, "01-2")))
    workbook = openpyxl.load_workbook(RAW / name, data_only=True)
    sheet = workbook["Ⅰ-9"]
    year_row = next(r for r in range(6, sheet.max_row) if sheet[f"A{r}"].value == year)
    for offset, (key, label) in enumerate([("total", "계"), *PROVIDERS], 2):
        address = f"{openpyxl.utils.get_column_letter(offset)}{year_row}"
        data["institutions_by_type"].append({"year": year, "provider_type": key, "institutions": integer(sheet, address), "unit": "institutions", "source": loc(name, sheet.title, address, year, 40)})
    assert sum(integer(sheet, f"{openpyxl.utils.get_column_letter(c)}{year_row}") for c in range(3, 19)) == integer(sheet, f"B{year_row}")
    for i, region in enumerate(REGIONS):
        row = year_row + 1 + i
        counts = {key: integer(sheet, f"{openpyxl.utils.get_column_letter(c)}{row}") for c, (key, _) in enumerate(PROVIDERS, 3)}
        total = integer(sheet, f"B{row}")
        assert sum(counts.values()) == total
        data["institutions_by_region_and_type"].append({"year": year, "region": region, "label_ko": "".join(sheet[f"A{row}"].value.split()), "total": total, "counts": counts, "unit": "institutions", "source": loc(name, sheet.title, f"B{row}:R{row}", year, 40)})
    for c in range(2, 19):
        letter = openpyxl.utils.get_column_letter(c)
        assert sum(integer(sheet, f"{letter}{r}") for r in range(year_row + 1, year_row + 18)) == integer(sheet, f"{letter}{year_row}")
    sheet = workbook["Ⅰ-13"]
    for key, row in [("total", 7), *[(key, i + 9) for i, (key, _) in enumerate(PROVIDERS)]]:
        metrics = {q: integer(sheet, f"{c}{row}") for q, c in zip(QUALIFICATIONS, "DEFGH")}
        assert metrics["physicians"] == sum(metrics[q] for q in QUALIFICATIONS[1:])
        data["physicians_by_type"].append({"year": year, "provider_type": key, **metrics, "unit": "persons", "source": loc(name, sheet.title, f"D{row}:H{row}", year, 48)})
    for c in "DEFGH":
        assert sum(integer(sheet, f"{c}{r}") for r in range(9, 25)) == integer(sheet, f"{c}7")
    sheet = workbook["Ⅰ-15"]
    for region, row in [("total", 7), *[(key, i + 8) for i, key in enumerate(REGIONS)]]:
        metrics = {q: integer(sheet, f"{c}{row}") for q, c in zip(QUALIFICATIONS, "CDEFG")}
        assert metrics["physicians"] == sum(metrics[q] for q in QUALIFICATIONS[1:])
        data["physicians_by_region"].append({"year": year, "region": region, **metrics, "unit": "persons", "source": loc(name, sheet.title, f"C{row}:G{row}", year, 56)})
    for c in "CDEFG":
        assert sum(integer(sheet, f"{c}{r}") for r in range(8, 25)) == integer(sheet, f"{c}7")
    if year == 2024:
        sheet = workbook["Ⅰ-10"]
        for key, row in [("total", 6), *[(key, i + 8) for i, (key, _) in enumerate(PROVIDERS)]]:
            counts = {form: integer(sheet, f"{openpyxl.utils.get_column_letter(c)}{row}") for c, form in enumerate(FORMS, 4)}
            assert sum(counts.values()) == integer(sheet, f"C{row}")
            data["institutions_by_establishment_2024"].append({"provider_type": key, "total": integer(sheet, f"C{row}"), "counts": counts, "unit": "institutions", "source": loc(name, sheet.title, f"C{row}:R{row}", year, 42)})

name = "nhis-care-2024.xlsx"
data["sources"].append(file_info(name, f"{BASE}?mode=download&articleNo=11007426&attachNo=365189", archive_file_id="nhis-yearbook-2024-tables.zip", archive_member=original_member("nhis-yearbook-2024-tables.zip", "03-1")))
sheet = openpyxl.load_workbook(RAW / name, data_only=True)["3-4"]
care_groups = [("total", 8), ("medical_and_community_health_institutions", 13), *[(key, 18 + 5 * i) for i, (key, _) in enumerate(PROVIDERS[:9])], ("community_health_agencies", 63), ("korean_medicine_hospital", 68), ("korean_medicine_clinic", 73)]
for key, start in care_groups:
    records = []
    for offset, setting in enumerate(["total", "inpatient", "outpatient"]):
        row = start + offset
        record = {"provider_type": key, "setting": setting, "patients": integer(sheet, f"D{row}"), "visit_days": integer(sheet, f"E{row}"), "reimbursed_days": integer(sheet, f"F{row}"), "medical_expenses_thousand_krw": str(Decimal(str(sheet[f"G{row}"].value))), "insurer_benefits_thousand_krw": str(Decimal(str(sheet[f"H{row}"].value))), "source": loc(name, sheet.title, f"D{row}:H{row}", 2024, 144)}
        records.append(record)
        data["care_by_type_2024"].append(record)
    for metric in ["visit_days", "reimbursed_days"]:
        assert records[0][metric] == records[1][metric] + records[2][metric]
    for metric in ["medical_expenses_thousand_krw", "insurer_benefits_thousand_krw"]:
        assert abs(Decimal(records[0][metric]) - Decimal(records[1][metric]) - Decimal(records[2][metric])) < Decimal("0.01")

data["validation"] = {
    "result": "passed",
    "checks": ["All extracted institution categories sum to national and regional totals in both years", "Every institution category sums across the 17 regions to its national value", "Every physician qualification split sums to physician count", "Institution and region physician margins each reconcile to national qualifications", "Every 2024 establishment-form split reconciles to institution count", "Care inpatient/outpatient days reconcile exactly; monetary subtotals reconcile within 0.01 thousand KRW", "2024 provider/care PDF table headers and footnotes visually inspected; workbook counts agree with PDF", "NHIS 2024 PDF SHA-256 equals independently downloaded SMHDB mirror"],
    "not_established": ["HIRA corrected attachment equality or corrigendum applicability", "2026 population levels", "Physician region × institution × employment-role joint distribution", "Clinical FTE from headcount", "Within-group physician income distributions"],
}
OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n")
print(f"Wrote {OUT}: " + ", ".join(f"{key}={len(value)}" for key, value in data.items() if isinstance(value, list)))
